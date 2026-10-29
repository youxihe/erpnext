# -*- coding: utf-8 -*-
"""cn_mfg_poc.bot_write — 企业微信 AI 机器人的受控写入（BUILD-60）。

沿用旧 AIERP POC 的「草稿 → 本人明确确认 → 过账」SOP，但完全落在 ERPNext 原生单据上：

- 预览/草稿阶段**不改变任何库存**：采购入库落在 Purchase Intake Draft，客户报备落在 Bot Draft；
- 确认只能由本人发起（适配器侧把「确认采购xxx / 确认报备xxx」做成 before_dispatch 拦截，
  **不提供 confirm 工具给模型**，避免模型替用户确认）；
- 幂等：文件指纹（sha256）/ 幂等键相同 → 返回原草稿；重复确认采购 → 返回原入库单，不重复过账；
- 取消：草稿置 Cancelled，不发生任何库存变化；
- 权限与 PC/H5 完全一致：确认时调用 cn_mfg_poc.api 的原生函数（角色 + 仓库 + 供应商校验）。
"""

import base64
import binascii
import io
import json
import re

import frappe
from frappe import _

from cn_mfg_poc.api import PURCHASE_ROLES, _require_any_role, cancel_purchase_draft
from cn_mfg_poc.bot import _log, _require, _serve

MODULE = "cn_mfg_poc.bot_write"
SALES_ROLES = {"Sales User", "Sales Manager", "System Manager"}
MAX_FILE_BYTES = 2 * 1024 * 1024

# Excel 模板列（中文表头）
TEMPLATE_COLUMNS = ["物料编码", "物料名称", "数量", "单位", "单价", "供应商", "仓库"]
HEADER_ALIASES = {
    "物料编码": "item_code",
    "物料code": "item_code",
    "item_code": "item_code",
    "编码": "item_code",
    "物料名称": "item_name",
    "名称": "item_name",
    "数量": "qty",
    "qty": "qty",
    "单位": "uom",
    "单价": "rate",
    "价格": "rate",
    "供应商": "supplier",
    "仓库": "warehouse",
}


def _clean(v):
    return str(v).strip() if v is not None else ""


def _num(v):
    if v is None or _clean(v) == "":
        return None
    try:
        return float(str(v).strip())
    except Exception:
        return None


# ---------------------------------------------------------------- 1. Excel 模板


def _template(_params):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "采购入库"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(["SENSOR", "传感器", 10, "件", 300, "", "原材料仓 - QC"])
    ws.append(["FASTENER", "紧固件", 20, "件", 5, "", "原材料仓 - QC"])
    buf = io.BytesIO()
    wb.save(buf)
    return {
        "file_name": "采购入库模板.xlsx",
        "file_base64": base64.b64encode(buf.getvalue()).decode(),
        "columns": TEMPLATE_COLUMNS,
        "hint": _("第一行是表头，请勿修改列名；供应商/仓库可留空，届时由对话补充。"),
    }


# ---------------------------------------------------------------- 2. Excel 解析 → 草稿


def _read_rows(file_base64, file_name):
    from openpyxl import load_workbook

    try:
        raw = base64.b64decode(file_base64)
    except (binascii.Error, ValueError):
        frappe.throw(_("文件不是合法的 base64 内容"), frappe.ValidationError)
    if not raw:
        frappe.throw(_("文件内容为空"), frappe.ValidationError)
    if len(raw) > MAX_FILE_BYTES:
        frappe.throw(_("文件超过 {0}MB，请精简后重发").format(MAX_FILE_BYTES // 1024 // 1024), frappe.ValidationError)
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_idx, headers = None, {}
    for i, row in enumerate(rows[:10]):
        if not row:
            continue
        mapped = {}
        ok = 0
        for cell in row:
            key = HEADER_ALIASES.get(_clean(cell))
            if key:
                ok += 1
        if ok >= 3:
            header_idx = i
            headers = {
                j: HEADER_ALIASES.get(_clean(cell)) for j, cell in enumerate(row) if HEADER_ALIASES.get(_clean(cell))
            }
            break
    if header_idx is None:
        frappe.throw(_("未识别到表头，请使用模板（列名：{0}）").format("、".join(TEMPLATE_COLUMNS)), frappe.ValidationError)
    out = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row or all(c is None or _clean(c) == "" for c in row):
            continue
        rec = {"row_no": offset}
        for j, key in headers.items():
            if j < len(row):
                rec[key] = row[j]
        out.append(rec)
    if not out:
        frappe.throw(_("表头之后没有数据行"), frappe.ValidationError)
    return out


def _parse_purchase_excel(p):
    from cn_mfg_poc import api

    file_sha256 = _clean(p.get("file_sha256"))
    file_name = _clean(p.get("file_name")) or "upload.xlsx"
    default_supplier = _clean(p.get("supplier"))
    default_warehouse = _clean(p.get("warehouse")) or "原材料仓 - QC"
    raw_rows = _read_rows(p.get("file_base64"), file_name)

    errors, lines, suppliers, warehouses = [], [], set(), set()
    for r in raw_rows:
        item = _clean(r.get("item_code"))
        qty = _num(r.get("qty"))
        rate = _num(r.get("rate"))
        supplier = _clean(r.get("supplier")) or default_supplier
        warehouse = _clean(r.get("warehouse")) or default_warehouse
        if not item:
            errors.append(_("第 {0} 行：物料编码为空").format(r["row_no"]))
            continue
        if not frappe.db.exists("Item", {"name": item, "disabled": 0}):
            errors.append(_("第 {0} 行：物料不存在或已停用 {1}").format(r["row_no"], item))
            continue
        if qty is None or qty <= 0:
            errors.append(_("第 {0} 行：数量必须是大于 0 的数字（当前 {1}）").format(r["row_no"], _clean(r.get("qty"))))
            continue
        if rate is None:
            rate = float(frappe.db.get_value("Item", item, "standard_rate") or 0)
        if rate < 0:
            errors.append(_("第 {0} 行：单价不能为负").format(r["row_no"]))
            continue
        if not frappe.db.exists("Supplier", {"name": supplier, "disabled": 0}):
            errors.append(_("第 {0} 行：供应商不存在或已停用 {1}").format(r["row_no"], supplier or "(空)"))
            continue
        if not frappe.db.exists("Warehouse", {"name": warehouse, "company": "昆山启程智能设备有限公司", "is_group": 0}):
            errors.append(_("第 {0} 行：入库仓库无效或不属于本公司 {1}").format(r["row_no"], warehouse))
            continue
        suppliers.add(supplier)
        warehouses.add(warehouse)
        lines.append({"item_code": item, "qty": qty, "rate": rate})

    if errors:
        return {
            "ok_preview": False,
            "errors": errors[:10],
            "error_count": len(errors),
            "hint": _("文件未通过校验，未创建草稿，库存没有任何变化。请修改后重新发送。"),
        }
    if len(suppliers) > 1:
        return {
            "ok_preview": False,
            "errors": [_("一份文件只能有一个供应商，当前：{0}").format("、".join(sorted(suppliers)))],
            "error_count": 1,
            "hint": _("请按供应商拆分文件后重发。"),
        }
    if len(warehouses) > 1:
        return {
            "ok_preview": False,
            "errors": [_("一份文件只能有一个入库仓库，当前：{0}").format("、".join(sorted(warehouses)))],
            "error_count": 1,
            "hint": _("请按仓库拆分文件后重发。"),
        }

    supplier = sorted(suppliers)[0]
    warehouse = sorted(warehouses)[0]
    res = api.prepare_purchase_draft(
        supplier=supplier,
        lines=lines,
        warehouse=warehouse,
        source="Bot-Excel",
        file_name=file_name,
        file_sha256=file_sha256 or None,
        wecom_userid=_clean(p.get("wecom_userid")),
    )
    draft = res["draft"]
    rows = []
    total_qty = total_amount = 0.0
    for row in res["lines"]:
        total_qty += row["qty"]
        total_amount += row["qty"] * (row.get("rate") or 0)
        rows.append(
            {
                "item_code": row["item_code"],
                "item_name": frappe.db.get_value("Item", row["item_code"], "item_name"),
                "uom": frappe.db.get_value("Item", row["item_code"], "stock_uom"),
                "qty": row["qty"],
                "rate": row.get("rate"),
            }
        )
    return {
        "ok_preview": True,
        "draft": draft,
        "replayed": res.get("replayed", False),
        "supplier": supplier,
        "warehouse": warehouse,
        "rows": rows,
        "total_qty": total_qty,
        "total_amount": round(total_amount, 2),
        "stock_hint": _("尚未入库，库存不变；确认后才会生成正式采购入库单并增加库存。"),
        "confirm_phrase": f"确认采购 {draft}",
    }


# ---------------------------------------------------------------- 3. 客户报备草稿


def _prepare_customer_report(p):
    customer_name = _clean(p.get("customer_name"))
    contact_name = _clean(p.get("contact_name"))
    mobile_no = _clean(p.get("mobile_no"))
    note = _clean(p.get("note"))
    idempotency_key = _clean(p.get("idempotency_key"))
    _require(customer_name, _("客户名称不能为空"), frappe.ValidationError)
    _require_any_role(SALES_ROLES, "报备客户")

    cust = frappe.db.get_value("Customer", {"customer_name": customer_name}, "name")
    if cust:
        return {"ok_preview": False, "existing": "customer", "customer": cust,
                "hint": _("该客户已存在（{0}），无需报备").format(cust)}
    dup = frappe.db.get_value(
        "Lead",
        {"company_name": customer_name, "status": ("not in", ["Converted", "Lost Quotation", "Do Not Contact"])},
        "name",
    )
    if dup:
        return {"ok_preview": False, "existing": "lead", "lead": dup,
                "hint": _("已有未关闭的报备单（{0}），请勿重复报备").format(dup)}

    if idempotency_key:
        old = frappe.db.get_value(
            "Bot Draft",
            {"idempotency_key": idempotency_key, "kind": "客户报备", "status": "Draft", "owner": frappe.session.user},
            "name",
        )
        if old:
            return {"ok_preview": True, "draft": old, "replayed": True,
                    "confirm_phrase": f"确认报备 {old}",
                    "hint": _("该报备已生成过待确认单，未重复创建。")}

    payload = {"customer_name": customer_name, "contact_name": contact_name,
               "mobile_no": mobile_no, "note": note}
    d = frappe.new_doc("Bot Draft")
    d.kind = "客户报备"
    d.wecom_userid = _clean(p.get("wecom_userid"))
    d.erp_user = frappe.session.user
    d.summary = f"{customer_name} · 联系人 {contact_name or '-'} · {mobile_no or '-'}"
    d.payload_json = json.dumps(payload, ensure_ascii=False)
    d.idempotency_key = idempotency_key or None
    d.status = "Draft"
    d.insert()
    frappe.db.commit()
    return {
        "ok_preview": True,
        "draft": d.name,
        "replayed": False,
        "customer_name": customer_name,
        "contact_name": contact_name,
        "mobile_no": mobile_no,
        "stock_hint": _("尚未创建报备单，确认后才会写入 ERP。"),
        "confirm_phrase": f"确认报备 {d.name}",
    }


# ---------------------------------------------------------------- 4. 确认 / 取消 / 待确认列表


def _resolve_draft(kind, draft):
    """按单据号或数字后缀解析草稿，并校验归属（本人或系统管理员）。"""
    name = _clean(draft)
    if re.fullmatch(r"\d+", name):
        prefix = "PID" if kind == "PURCHASE" else "BD"
        name = f"{prefix}-{int(name):05d}"
    doctype = "Purchase Intake Draft" if kind == "PURCHASE" else "Bot Draft"
    if not frappe.db.exists(doctype, name):
        frappe.throw(_("找不到待确认单 {0}").format(name), frappe.ValidationError)
    d = frappe.get_doc(doctype, name)
    if d.owner != frappe.session.user and "System Manager" not in frappe.get_roles():
        frappe.throw(_("该待确认单不属于您（{0}）").format(d.owner), frappe.PermissionError)
    return d


def _confirm_draft(p):
    from cn_mfg_poc import api

    kind = (_clean(p.get("kind")) or "").upper()
    _require(kind in ("PURCHASE", "CUSTOMER"), _("确认类型必须是 PURCHASE 或 CUSTOMER"), frappe.ValidationError)
    d = _resolve_draft(kind, p.get("draft"))

    if kind == "PURCHASE":
        res = api.confirm_purchase_draft(d.name)
        d.reload()
        _log("bot_confirm", tool="confirm_draft", kind=kind, draft=d.name,
             purchase_receipt=res.get("purchase_receipt"), already=res.get("already_confirmed"))
        return {
            "kind": "PURCHASE",
            "draft": d.name,
            "status": d.status,
            "purchase_receipt": res.get("purchase_receipt"),
            "docstatus": res.get("docstatus"),
            "already_confirmed": bool(res.get("already_confirmed")),
            "message": res.get("message")
            or _("已生成并提交采购入库单 {0}，库存已按单增加一次。").format(res.get("purchase_receipt")),
        }

    # 客户报备
    if d.status == "Confirmed" and d.result_name:
        return {"kind": "CUSTOMER", "draft": d.name, "status": "Confirmed", "lead": d.result_name,
                "already_confirmed": True, "message": _("该报备已确认过，返回原单据，未重复创建。")}
    _require(d.status == "Draft", _("待确认单状态为 {0}，不能确认").format(d.status), frappe.ValidationError)
    payload = json.loads(d.payload_json)
    res = api.report_customer(
        customer_name=payload.get("customer_name"),
        contact_name=payload.get("contact_name"),
        mobile_no=payload.get("mobile_no"),
        note=payload.get("note"),
    )
    if res.get("existing"):
        d.status = "Cancelled"
        d.summary = (d.summary or "") + f" · 已存在{res.get('existing')}，取消"
        d.flags.ignore_permissions = True
        d.save()
        frappe.db.commit()
        return {"kind": "CUSTOMER", "draft": d.name, "status": "Cancelled",
                "existing": res.get("existing"), "message": res.get("message")}
    d.status = "Confirmed"
    d.result_doctype = "Lead"
    d.result_name = res.get("lead")
    d.flags.ignore_permissions = True
    d.save()
    frappe.db.commit()
    _log("bot_confirm", tool="confirm_draft", kind=kind, draft=d.name, lead=res.get("lead"))
    return {
        "kind": "CUSTOMER",
        "draft": d.name,
        "status": "Confirmed",
        "lead": res.get("lead"),
        "lead_owner": res.get("owner"),
        "already_confirmed": False,
        "message": _("已创建客户报备 {0}，PC 与 H5 立即可见。").format(res.get("lead")),
    }


def _cancel_draft(p):
    kind = (_clean(p.get("kind")) or "").upper()
    _require(kind in ("PURCHASE", "CUSTOMER"), _("取消类型必须是 PURCHASE 或 CUSTOMER"), frappe.ValidationError)
    d = _resolve_draft(kind, p.get("draft"))
    if kind == "PURCHASE":
        res = cancel_purchase_draft(d.name)
        _log("bot_cancel", tool="cancel_draft", kind=kind, draft=d.name)
        return {"kind": "PURCHASE", "draft": d.name, "status": "Cancelled", "stock_unchanged": True,
                "message": _("草稿 {0} 已取消，未生成入库单，库存不变。").format(d.name)}
    _require(d.status == "Draft", _("待确认单状态为 {0}，不能取消").format(d.status), frappe.ValidationError)
    d.status = "Cancelled"
    d.flags.ignore_permissions = True
    d.save()
    frappe.db.commit()
    _log("bot_cancel", tool="cancel_draft", kind=kind, draft=d.name)
    return {"kind": "CUSTOMER", "draft": d.name, "status": "Cancelled", "stock_unchanged": True,
            "message": _("报备草稿 {0} 已取消，未创建任何 ERP 单据。").format(d.name)}


def _my_drafts(p):
    out = {"purchase": [], "customer": []}
    if set(frappe.get_roles()) & PURCHASE_ROLES:
        out["purchase"] = frappe.get_all(
            "Purchase Intake Draft",
            filters={"status": "Draft", "owner": frappe.session.user},
            fields=["name", "supplier", "warehouse", "creation", "file_name"],
            order_by="creation desc",
            limit=10,
        )
    if set(frappe.get_roles()) & SALES_ROLES:
        out["customer"] = frappe.get_all(
            "Bot Draft",
            filters={"kind": "客户报备", "status": "Draft", "owner": frappe.session.user},
            fields=["name", "summary", "creation"],
            order_by="creation desc",
            limit=10,
        )
    return out


# ---------------------------------------------------------------- 白名单端点


@frappe.whitelist(methods=["POST"])
def purchase_excel_template(**kwargs):
    """下载采购入库 Excel 模板（base64），不产生任何业务数据。"""
    return _serve("purchase_excel_template", _template, MODULE)


@frappe.whitelist(methods=["POST"])
def parse_purchase_excel(**kwargs):
    """解析并逐行校验 Excel → 生成 Purchase Intake Draft（草稿，不改库存）。"""
    return _serve("parse_purchase_excel", _parse_purchase_excel, MODULE)


@frappe.whitelist(methods=["POST"])
def prepare_customer_report(**kwargs):
    """自然语言客户报备 → 查重 → 生成 Bot Draft（待确认，不写 ERP）。"""
    return _serve("prepare_customer_report", _prepare_customer_report, MODULE)


@frappe.whitelist(methods=["POST"])
def confirm_draft(**kwargs):
    """本人明确确认后过账：采购→Purchase Receipt；报备→Lead。幂等。"""
    return _serve("confirm_draft", _confirm_draft, MODULE)


@frappe.whitelist(methods=["POST"])
def cancel_draft(**kwargs):
    """取消尚未过账的草稿，不发生库存变化。"""
    return _serve("cancel_draft", _cancel_draft, MODULE)


@frappe.whitelist(methods=["POST"])
def my_drafts(**kwargs):
    """列出当前用户待确认的草稿（采购入库 / 客户报备）。"""
    return _serve("my_drafts", _my_drafts, MODULE)
